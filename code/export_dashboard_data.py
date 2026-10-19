"""
Step 7 (part 1): Export the workbook into dashboard_data.json.

Reads every sheet of Loan_Collection_Prototype.xlsx (the synthetic data,
the rules engine's recommendations, the AI drafts, the approvals, and
the audit log) and writes it all out as one JSON file that the
dashboard page reads.

Run this after re-running rules_engine.py / ai_layer.py / approval_layer.py,
whenever the underlying data changes, then run build_dashboard.py to
bake the new JSON into the dashboard page.
"""

import json
from datetime import datetime, date

import openpyxl

WORKBOOK_PATH = "Loan_Collection_Prototype.xlsx"
OUTPUT_PATH = "dashboard_data.json"


def sheet_to_dicts(ws):
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        d = {}
        for h, v in zip(headers, row):
            if isinstance(v, (datetime, date)):
                v = v.isoformat()[:10]
            d[h] = v
        rows.append(d)
    return rows


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)

    data = {
        "loans": sheet_to_dicts(wb["Loan_Contracts"]),
        "cases": sheet_to_dicts(wb["Collection_Cases"]),
        "payments": sheet_to_dicts(wb["Payments"]),
        "bureau": sheet_to_dicts(wb["Bureau_Reporting_Log"]),
        "recommendations": sheet_to_dicts(wb["Engine_Recommendations"]),
        "ai_drafts": sheet_to_dicts(wb["AI_Drafts"]),
        "approvals": sheet_to_dicts(wb["Approval_Queue"]),
        "audit_log": sheet_to_dicts(wb["Audit_Log"]),
    }

    with open(OUTPUT_PATH, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=1, default=str)

    print(f"Wrote {OUTPUT_PATH}")
    for key, rows in data.items():
        print(f"  {key}: {len(rows)} rows")


if __name__ == "__main__":
    main()
